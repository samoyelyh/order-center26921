from __future__ import annotations

import io
import json
import mimetypes
import posixpath
import zipfile
from collections import defaultdict, deque
from collections.abc import Iterable
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import (
    OrderAsset,
    OrderBuyerAsset,
    OrderImportBatch,
    OrderImportBatchItem,
    OrderItem,
)
from app.orders.categories.blade_shoes import BladeShoesParser
from app.orders.core.category_resolver import CategoryResolver
from app.orders.core.generic_parser import GenericOrderParser
from app.orders.core.parser_registry import ParserRegistry
from app.services.fingerprint import compute_blake3
from app.services.repository import new_id, utcnow
from app.services.storage import StorageError, build_storage_key, get_storage


class OrderImportError(ValueError):
    pass


def default_registry() -> ParserRegistry:
    registry = ParserRegistry()
    registry.register("BLADE_SHOES", BladeShoesParser())
    return registry


def _mime(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".json"):
        return "application/json"
    if lower.endswith(".zip"):
        return "application/zip"
    if lower.endswith(".svg"):
        return "image/svg+xml"
    return mimetypes.guess_type(filename)[0] or "application/octet-stream"


def _ensure_asset(db: Session, filename: str, data: bytes, actor: str) -> OrderAsset:
    """把订单原始文件存为 OrderAsset（order-center 自有表，不用素材库 assets）。"""
    digest = compute_blake3(data)
    existing = db.execute(
        select(OrderAsset).where(OrderAsset.blake3 == digest)
    ).scalar_one_or_none()
    if existing is not None:
        return existing
    storage = get_storage()
    storage_key = build_storage_key(filename)
    try:
        storage.put(storage_key, data, _mime(filename))
    except StorageError as exc:
        raise OrderImportError(f"原始文件写入存储失败：{filename}") from exc
    asset = OrderAsset(
        id=new_id("order-asset"),
        storage_key=storage_key,
        original_filename=filename,
        mime_type=_mime(filename),
        size_bytes=len(data),
        blake3=digest,
        created_by=actor,
    )
    db.add(asset)
    db.flush()
    return asset


def _safe_archive_files(
    data: bytes,
    filename: str,
    *,
    prefix: str = "",
    depth: int = 0,
) -> list[tuple[str, bytes]]:
    """Read nested ZIP members in memory with traversal/depth/size guards."""
    if depth > 5:
        raise OrderImportError("订单 ZIP 嵌套层级超过限制")
    if not filename.lower().endswith(".zip"):
        return [(prefix or filename, data)]
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            out: list[tuple[str, bytes]] = []
            for info in archive.infolist()[:5000]:
                if info.is_dir() or info.file_size > 50 * 1024 * 1024:
                    continue
                name = info.filename.replace("\\", "/")
                if name.startswith("/") or ".." in name.split("/"):
                    continue
                content = archive.read(info)
                logical_path = f"{prefix}{name}"
                if name.lower().endswith(".zip"):
                    # Keep the container itself as an OrderAsset source while
                    # also expanding it for parsing.
                    out.append((logical_path, content))
                    out.extend(
                        _safe_archive_files(
                            content,
                            name,
                            prefix=f"{logical_path}!/",
                            depth=depth + 1,
                        )
                    )
                else:
                    out.append((logical_path, content))
            return out
    except zipfile.BadZipFile as exc:
        raise OrderImportError(f"订单文件不是有效 ZIP：{filename}") from exc


def _safe_json_files(data: bytes, filename: str) -> list[tuple[str, bytes]]:
    return [
        (path, content)
        for path, content in _safe_archive_files(data, filename)
        if path.lower().endswith(".json")
    ]


def _dedupe_key(record, raw_path: str = "") -> str:
    if record.order_item_id:
        return f"ITEM:{record.order_item_id.strip()}"
    # 路径不是订单身份：同一订单重新导出时 ZIP 内部路径可能变化。缺少
    # orderItemId 时使用稳定业务字段 + 已提取定制内容，避免重复导入重复计销量。
    stable = json.dumps(
        {
            "orderId": record.order_id,
            "asin": record.child_asin,
            "sku": record.sku,
            "quantity": record.quantity,
            "images": [
                [c.thumbnail_url, c.overlay_url, c.image_name, c.label]
                for c in record.image_candidates
            ],
            "texts": [[f.label, f.value] for f in record.text_fields],
            "options": record.option_values,
            "buyerOriginal": record.buyer_logo_original,
            "buyerSvg": record.buyer_logo_svg,
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return f"FALLBACK:{compute_blake3(stable)}"


def _json_documents(files: Iterable[tuple[str, bytes]]) -> Iterable[tuple[str, bytes, Any]]:
    for path, raw in files:
        if not path.lower().endswith(".json"):
            continue
        try:
            document = json.loads(raw.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            continue
        yield path, raw, document


def _excel_rows(files: list[tuple[str, bytes]]) -> dict[tuple[str, str], deque[dict[str, Any]]]:
    """Parse LingXing's outer workbook and index rows by (orderId, ASIN).

    The real export has one workbook row per order line and one nested customization
    ZIP per customized line.  The JSON carries orderId/orderItemId/ASIN but commonly
    lacks MSKU, while the workbook carries MSKU.  Pair multiplicity is preserved with
    a queue so multi-line orders remain one-to-one.
    """
    indexed: dict[tuple[str, str], deque[dict[str, Any]]] = defaultdict(deque)
    workbook_entry = next((pair for pair in files if pair[0].lower().endswith(".xlsx")), None)
    if workbook_entry is None:
        return indexed
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(
            io.BytesIO(workbook_entry[1]), read_only=False, data_only=True
        )
        sheet = workbook.active
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        columns = {name: index for index, name in enumerate(headers)}
        required = {"平台单号", "ASIN/商品Id"}
        if not required.issubset(columns):
            return indexed
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            order_id = str(values[columns["平台单号"]] or "").strip()
            asin = str(values[columns["ASIN/商品Id"]] or "").strip()
            if not order_id or not asin:
                continue
            image_cell = sheet.cell(row=row_number, column=columns.get("图片", -1) + 1) if "图片" in columns else None
            image_path = image_cell.hyperlink.target if image_cell and image_cell.hyperlink else None
            indexed[(order_id, asin)].append({
                "sku": str(values[columns["MSKU"]]).strip() if "MSKU" in columns and values[columns["MSKU"]] else None,
                "quantity": int(values[columns["数量"]] or 1) if "数量" in columns else 1,
                "buyerRequest": str(values[columns["买家留言"]]).strip() if "买家留言" in columns and values[columns["买家留言"]] else None,
                "productRemark": str(values[columns["商品备注"]]).strip() if "商品备注" in columns and values[columns["商品备注"]] else None,
                "orderImagePath": image_path,
                "rowNumber": row_number,
            })
    except Exception as exc:  # noqa: BLE001
        raise OrderImportError("外层订单 Excel 解析失败") from exc
    return indexed


def _buyer_file(
    files: list[tuple[str, bytes]], entry: dict[str, str | None], json_path: str
) -> tuple[str, bytes] | None:
    """Resolve imageName (preferred) or URL basename to a file in the same export."""
    names: list[str] = []
    for value in (entry.get("name"), entry.get("url")):
        if not value:
            continue
        clean = value.split("?", 1)[0].replace("\\", "/")
        base = posixpath.basename(clean).strip().lower()
        if base:
            names.append(base)
    if not names:
        return None
    prefix = json_path.rsplit("!/", 1)[0] + "!/" if "!/" in json_path else ""
    matches = [
        pair for pair in files
        if pair[0].startswith(prefix)
        and posixpath.basename(pair[0]).lower() in names
    ]
    return matches[0] if len(matches) == 1 else None


def import_order_zip(
    db: Session,
    *,
    filename: str,
    data: bytes,
    actor: str,
    category_code: str | None = None,
    registry: ParserRegistry | None = None,
) -> tuple[OrderImportBatch, dict[str, int]]:
    if not data:
        raise OrderImportError("订单 ZIP 不能为空")
    if not filename.lower().endswith(".zip"):
        raise OrderImportError("订单导入只接受 ZIP 文件")
    registry = registry or default_registry()
    resolver = CategoryResolver()
    batch_category = resolver.resolve(manual_code=category_code)
    zip_hash = compute_blake3(data)
    raw_zip = _ensure_asset(db, filename, data, actor)
    canonical = db.execute(
        select(OrderImportBatch)
        .where(OrderImportBatch.zip_blake3 == zip_hash)
        .order_by(OrderImportBatch.created_at.asc())
        .limit(1)
    ).scalar_one_or_none()
    batch = OrderImportBatch(
        id=new_id("order-import"),
        original_filename=filename,
        raw_zip_asset_id=raw_zip.id,
        zip_blake3=zip_hash,
        category_code=batch_category.code,
        category_name=batch_category.name,
        category_source=batch_category.source,
        category_confidence=batch_category.confidence,
        status="DUPLICATE" if canonical is not None else "PARSED",
        duplicate_of_batch_id=canonical.id if canonical is not None else None,
        created_by=actor,
    )
    db.add(batch)
    db.flush()
    if canonical is not None:
        batch.total_json_count = canonical.total_json_count
        batch.total_item_count = canonical.total_item_count
        db.commit()
        return batch, _summary_stats(db, canonical.id, {"newItemCount": 0, "duplicateItemCount": batch.total_item_count})

    archive_files = _safe_archive_files(data, filename)
    for source_path, source_bytes in archive_files:
        if source_path.lower().endswith((".xlsx", ".zip")):
            _ensure_asset(db, source_path, source_bytes, actor)
    excel_rows = _excel_rows(archive_files)
    ambiguous_excel_keys = {
        key for key, rows in excel_rows.items()
        if len({(row.get("sku"), row.get("quantity")) for row in rows}) > 1
    }
    has_excel = any(path.lower().endswith(".xlsx") for path, _ in archive_files)
    json_count = item_count = new_count = duplicate_count = 0
    review_count = 0
    seen_item_ids: set[str] = set()
    fallback_occurrences: dict[str, int] = defaultdict(int)
    noncustom_occurrences: dict[str, int] = defaultdict(int)
    for json_path, raw_json, document in _json_documents(archive_files):
        json_count += 1
        json_asset = _ensure_asset(db, json_path, raw_json, actor)
        records = GenericOrderParser().parse(document, source_path=json_path)
        for record in records:
            item_count += 1
            excel_row = None
            excel_key = (record.order_id, record.child_asin or "")
            if excel_key not in ambiguous_excel_keys and excel_rows.get(excel_key):
                excel_row = excel_rows[excel_key].popleft()
                record.sku = record.sku or excel_row.get("sku")
                record.quantity = excel_row.get("quantity") or record.quantity
            category = resolver.resolve(manual_code=category_code, sku=record.sku, asin=record.child_asin)
            normalized = registry.parse(record, category=category)
            if excel_row is None and has_excel:
                normalized.parse_status = "REVIEW_REQUIRED"
                normalized.details["reason"] = (
                    "AMBIGUOUS_OUTER_EXCEL_MATCH"
                    if excel_key in ambiguous_excel_keys else "MISSING_OUTER_EXCEL_MATCH"
                )
            if excel_row is not None:
                normalized.buyer_request = normalized.buyer_request or excel_row.get("buyerRequest")
                normalized.details["outerExcel"] = {
                    "rowNumber": excel_row.get("rowNumber"),
                    "orderImagePath": excel_row.get("orderImagePath"),
                    "productRemark": excel_row.get("productRemark"),
                }
            if normalized.parse_status == "REVIEW_REQUIRED":
                review_count += 1
            # Parsing and material matching are independent workflows. Import
            # never calls the matcher and therefore starts in NOT_STARTED.
            normalized.match_status = "NOT_STARTED"
            key = _dedupe_key(record, record.raw_path)
            if not record.order_item_id:
                fallback_occurrences[key] += 1
                if fallback_occurrences[key] > 1:
                    key = f"{key}:{fallback_occurrences[key]}"
            item = db.execute(select(OrderItem).where(OrderItem.dedupe_key == key)).scalar_one_or_none()
            if item is None:
                item = OrderItem(
                    id=new_id("order-item"),
                    dedupe_key=key,
                    first_import_batch_id=batch.id,
                    order_id=normalized.order_id,
                    order_item_id=normalized.order_item_id,
                    child_asin=normalized.child_asin,
                    sku=normalized.sku,
                    quantity=normalized.quantity,
                    category_code=normalized.category_code,
                    category_name=normalized.category_name,
                    category_source=normalized.category_source,
                    category_confidence=normalized.category_confidence,
                    parser_version=normalized.parser_version,
                    parse_status=normalized.parse_status,
                    raw_json_hash=compute_blake3(raw_json),
                    normalized_payload=normalized.as_payload(),
                    match_status=normalized.match_status,
                )
                db.add(item)
                db.flush()
                new_count += 1
                has_original = bool(normalized.buyer_logo_original)
                for role, entries in (("BUYER_LOGO_ORIGINAL", normalized.buyer_logo_original), ("BUYER_LOGO_SVG", normalized.buyer_logo_svg)):
                    for index, entry in enumerate(entries):
                        resolved = _buyer_file(archive_files, entry, json_path)
                        buyer_asset = (
                            _ensure_asset(db, resolved[0], resolved[1], actor)
                            if resolved is not None
                            else None
                        )
                        db.add(OrderBuyerAsset(
                            id=new_id("buyer-asset"), order_item_id=item.id, role=role,
                            asset_id=buyer_asset.id if buyer_asset else None,
                            source_name=entry.get("name"), source_url=entry.get("url"),
                            source_path=entry.get("path"),
                            is_primary=index == 0 and (role == "BUYER_LOGO_ORIGINAL" or not has_original),
                        ))
            else:
                duplicate_count += 1
            if item.id not in seen_item_ids:
                db.add(OrderImportBatchItem(
                    id=new_id("order-import-item"), import_batch_id=batch.id, order_item_id=item.id,
                    raw_json_asset_id=json_asset.id, raw_json_path=record.raw_path,
                    raw_json_hash=compute_blake3(raw_json), parser_version=normalized.parser_version,
                    parse_status=normalized.parse_status, normalized_payload=normalized.as_payload(),
                ))
                seen_item_ids.add(item.id)
    # Keep workbook lines that have no customization JSON. They are valid
    # non-custom orders and must remain visible for production review.
    if has_excel:
        for (order_id, asin), rows in excel_rows.items():
            while rows:
                row = rows.popleft()
                sku = row.get("sku")
                category = resolver.resolve(manual_code=category_code, sku=sku, asin=asin)
                payload = {
                    "orderId": order_id, "orderItemId": None, "childAsin": asin,
                    "sku": sku, "quantity": row.get("quantity") or 1,
                    "categoryCode": category.code, "categoryName": category.name,
                    "categorySource": category.source, "categoryConfidence": category.confidence,
                    "parserVersion": "NO_JSON", "customizationStatus": "NON_CUSTOM",
                    "imageType": None, "parseStatus": "NON_CUSTOM",
                    "matchStatus": "NOT_STARTED", "details": {"outerExcel": row, "reason": "NO_CUSTOMIZATION_JSON"},
                }
                stable = json.dumps({
                    "orderId": order_id, "asin": asin, "sku": sku,
                    "quantity": row.get("quantity") or 1,
                    "buyerRequest": row.get("buyerRequest"),
                    "productRemark": row.get("productRemark"),
                }, sort_keys=True, ensure_ascii=False).encode()
                base_key = f"FALLBACK:{compute_blake3(stable)}"
                noncustom_occurrences[base_key] += 1
                key = base_key if noncustom_occurrences[base_key] == 1 else f"{base_key}:{noncustom_occurrences[base_key]}"
                if db.execute(select(OrderItem).where(OrderItem.dedupe_key == key)).scalar_one_or_none() is not None:
                    duplicate_count += 1
                    continue
                item = OrderItem(
                    id=new_id("order-item"), dedupe_key=key, first_import_batch_id=batch.id,
                    order_id=order_id, order_item_id=None, child_asin=asin, sku=sku,
                    quantity=row.get("quantity") or 1, category_code=category.code,
                    category_name=category.name, category_source=category.source,
                    category_confidence=category.confidence, parser_version="NO_JSON",
                    parse_status="NON_CUSTOM", raw_json_hash=compute_blake3(stable),
                    normalized_payload=payload, match_status="NOT_STARTED",
                )
                db.add(item)
                db.flush()
                item_count += 1
                new_count += 1
    batch.total_json_count = json_count
    batch.total_item_count = item_count
    if json_count == 0:
        batch.status = "FAILED"
    elif review_count:
        batch.status = "PARTIAL"
    db.commit()
    return batch, _summary_stats(db, batch.id, {"newItemCount": new_count, "duplicateItemCount": duplicate_count})


def _summary_stats(db: Session, batch_id: str, base: dict[str, int]) -> dict[str, int]:
    """Return import and parse counts for the upload confirmation UI."""
    items = db.execute(
        select(OrderItem).join(OrderImportBatchItem, OrderImportBatchItem.order_item_id == OrderItem.id)
        .where(OrderImportBatchItem.import_batch_id == batch_id)
    ).scalars().all()
    stats = {"jsonCount": sum(1 for _ in db.execute(select(OrderImportBatchItem.raw_json_asset_id).where(OrderImportBatchItem.import_batch_id == batch_id)).all()), "itemCount": len(items), **base}
    for status in ("PARSED", "REVIEW_REQUIRED", "FAILED", "NON_CUSTOM"):
        stats[status.lower()] = sum(1 for item in items if item.parse_status == status)
    for image_type in ("MATERIAL_SOURCE", "FINAL_EFFECT", "PREVIEW_ONLY", "UNKNOWN"):
        stats[image_type.lower()] = sum(1 for item in items if item.parse_status != "NON_CUSTOM" and (item.normalized_payload or {}).get("imageType") == image_type)
    stats["black"] = sum(1 for item in items if item.parse_status != "NON_CUSTOM" and (item.normalized_payload or {}).get("soleColor") == "BLACK")
    stats["white"] = sum(1 for item in items if item.parse_status != "NON_CUSTOM" and (item.normalized_payload or {}).get("soleColor") == "WHITE")
    stats["customName"] = sum(1 for item in items if (item.normalized_payload or {}).get("customName"))
    stats["customNumber"] = sum(1 for item in items if (item.normalized_payload or {}).get("customNumber"))
    stats["buyerRequest"] = sum(1 for item in items if (item.normalized_payload or {}).get("buyerRequest"))
    stats["buyerLogo"] = sum(1 for item in items if (item.normalized_payload or {}).get("hasBuyerLogo"))
    return stats
